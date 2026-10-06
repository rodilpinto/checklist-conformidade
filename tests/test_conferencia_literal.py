# -*- coding: utf-8 -*-
"""Conferência do texto literal contra o normativo, com rodadas de correção
(Rodrigo, 06/10/2026: "texto tem de ser literal. se der divergencia, precisamos de
mais rodadas até chegar em convergência total"). Sem rede: o corretor é simulado.
"""
import sys
from io import BytesIO
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from lib.conferencia import (  # noqa: E402
    NAO_CONFERE,
    CONFERE,
    confere,
    convergir,
    normalizar,
)
from lib.excel_builder import build_excel  # noqa: E402
from lib.llm import LLMError, validate_items  # noqa: E402

FONTE = (
    "Art. 6º O uso de dados pessoais por sistemas de IA deve ser evitado.\n"
    "Art. 7º Quando o tratamento de dados pessoais for estritamente\n"
    "necessário, deve observar a Lei nº 13.709, de 2018 , e as normas internas:\n"
    "I – registrar a “base legal”;\n"
)


def _itens(*textos):
    return validate_items([
        {"capitulo": "III", "artigo": f"Art. {i}", "texto_literal": t,
         "probabilidade": 3, "impacto": 3, "atores": [{"nome": "Ditec", "papel": "responsavel"}]}
        for i, t in enumerate(textos, start=1)
    ])


# ── confere / normalizar ──────────────────────────────────────────────────────

def test_confere_tolera_so_espacos_quebras_aspas_e_travessoes():
    fn = normalizar(FONTE)
    assert confere("Art. 7º Quando o tratamento de dados pessoais for estritamente necessário,", fn)
    assert confere('I - registrar a "base legal";', fn)
    assert confere("Lei nº 13.709, de 2018, e as normas internas", fn)  # espaço antes da vírgula


def test_nao_confere_com_acento_palavra_ou_pontuacao_diferente():
    fn = normalizar(FONTE)
    assert not confere("Art. 7º Quando o tratamento de dados pessoais for estritamente necessario,", fn)
    assert not confere("O uso de dados pessoais por sistemas de IA deve ser evitada.", fn)
    assert not confere("O uso de dados pessoais, por sistemas de IA", fn)
    assert not confere("", fn)
    assert not confere(None, fn)


# ── convergir ─────────────────────────────────────────────────────────────────

def test_tudo_confere_nao_chama_o_corretor():
    itens = _itens("O uso de dados pessoais por sistemas de IA deve ser evitado.")
    chamadas = []
    res = convergir(itens, FONTE, lambda d: chamadas.append(d) or {})
    assert chamadas == []
    assert res == {"total": 1, "rodadas": 0, "corrigidos": 0, "restantes": 0, "erro": None}
    assert itens[0]["conferencia_literal"] == CONFERE


def test_rodada_corrige_so_os_divergentes():
    itens = _itens("O uso de dados pessoais por sistemas de IA deve ser evitado.",
                   "O uso de dados pessoais deve ser evitado.")  # resumido
    pedidos = []

    def corrigir(divergentes):
        pedidos.append([it["id"] for it in divergentes])
        return {2: "O uso de dados pessoais por sistemas de IA deve ser evitado."}

    res = convergir(itens, FONTE, corrigir)
    assert pedidos == [[2]]
    assert res["rodadas"] == 1 and res["corrigidos"] == 1 and res["restantes"] == 0
    assert itens[1]["texto_literal"] == "O uso de dados pessoais por sistemas de IA deve ser evitado."
    assert [it["conferencia_literal"] for it in itens] == [CONFERE, CONFERE]


def test_varias_rodadas_ate_convergir():
    itens = _itens("texto inventado")
    respostas = iter([{1: "outro texto inventado"}, {1: "I – registrar a “base legal”;"}])
    res = convergir(itens, FONTE, lambda d: next(respostas))
    assert res["rodadas"] == 2 and res["restantes"] == 0
    assert itens[0]["texto_literal"] == "I – registrar a “base legal”;"


def test_limite_de_rodadas_marca_o_que_sobrou():
    itens = _itens("texto inventado")
    chamadas = []
    res = convergir(itens, FONTE, lambda d: chamadas.append(1) or {1: "ainda errado"}, max_rodadas=3)
    assert len(chamadas) == 3
    assert res["restantes"] == 1 and res["corrigidos"] == 0
    assert itens[0]["texto_literal"] == "texto inventado"  # resposta que não confere não substitui
    assert itens[0]["conferencia_literal"] == NAO_CONFERE


def test_falha_do_modelo_na_rodada_nao_derruba_a_geracao():
    itens = _itens("texto inventado")

    def corrigir(_):
        raise LLMError("Nenhum modelo de IA respondeu.")

    res = convergir(itens, FONTE, corrigir)
    assert res["restantes"] == 1 and "Nenhum modelo" in res["erro"]
    assert itens[0]["conferencia_literal"] == NAO_CONFERE


# ── planilha ──────────────────────────────────────────────────────────────────

def test_coluna_de_conferencia_ao_lado_do_texto_literal():
    itens = _itens("O uso de dados pessoais por sistemas de IA deve ser evitado.", "inventado")
    convergir(itens, FONTE, lambda d: {})
    ws = openpyxl.load_workbook(BytesIO(build_excel(itens))).worksheets[0]
    assert ws.cell(1, 4).value == "Texto Literal do Artigo"
    assert ws.cell(1, 5).value == "Conferência do\nTexto Literal"
    assert [ws.cell(r, 5).value for r in (2, 3)] == [CONFERE, NAO_CONFERE]
