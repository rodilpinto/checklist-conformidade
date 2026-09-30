# -*- coding: utf-8 -*-
"""Testes do extracao_texto. Viajam com a pasta: nao importam nada de app, nao fazem rede
(PDF e DOCX sao gerados em memoria; requests.get e socket.getaddrinfo sao dublados).

Rodar da pasta que CONTEM extracao_texto/:
    python -m pytest extracao_texto -q
"""
from __future__ import annotations

import io
import socket

import pytest

import extracao_texto
from extracao_texto import extractor as ex


# --- PDF / DOCX / texto ------------------------------------------------------------

def _pdf(*paginas: str) -> bytes:
    import fitz
    doc = fitz.open()
    for texto in paginas:
        doc.new_page().insert_text((72, 72), texto)
    dados = doc.tobytes()
    doc.close()
    return dados


def _docx(paragrafos, tabela=None) -> bytes:
    from docx import Document
    d = Document()
    for p in paragrafos:
        d.add_paragraph(p)
    if tabela:
        t = d.add_table(rows=len(tabela), cols=len(tabela[0]))
        for i, linha in enumerate(tabela):
            for j, celula in enumerate(linha):
                t.cell(i, j).text = celula
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_pdf_extrai_todas_as_paginas_com_marcador():
    texto = extracao_texto.extract_text(_pdf("Art. 1 Primeira pagina", "Art. 2 Segunda pagina"), "pdf")
    assert "--- Pagina 1 ---" in texto and "Art. 1 Primeira pagina" in texto
    assert "--- Pagina 2 ---" in texto and "Art. 2 Segunda pagina" in texto


def test_pdf_vazio_invalido_ou_grande_demais(monkeypatch):
    with pytest.raises(ValueError):
        ex.extract_from_pdf(b"")
    with pytest.raises(RuntimeError):
        ex.extract_from_pdf(b"isto nao e um pdf")
    monkeypatch.setattr(ex, "_MAX_FILE_SIZE_BYTES", 10)
    with pytest.raises(ValueError, match="excede o limite"):
        ex.extract_from_pdf(_pdf("x"))


def test_pdf_sem_texto_devolve_vazio():
    assert ex.extract_from_pdf(_pdf("")) == ""


def test_docx_extrai_paragrafos_e_tabelas():
    dados = _docx(["Art. 1 Objeto", "", "Art. 2 Prazo"], tabela=[["Requisito", "Prazo"], ["Relatorio", "30 dias"]])
    texto = extracao_texto.extract_text(dados, "docx")
    assert texto.splitlines() == ["Art. 1 Objeto", "Art. 2 Prazo", "Requisito | Prazo", "Relatorio | 30 dias"]


def test_texto_puro_e_truncamento(monkeypatch):
    assert extracao_texto.extract_text("  colado  ", "text") == "colado"
    monkeypatch.setattr(ex, "_MAX_EXTRACTED_CHARS", 10)
    assert len(extracao_texto.extract_text("a" * 50, "text")) == 10


@pytest.mark.parametrize("fonte,tipo", [("x", "pdf"), (b"x", "text"), (b"x", "url"), ("x", "xls")])
def test_tipo_errado_e_erro(fonte, tipo):
    with pytest.raises(ValueError):
        extracao_texto.extract_text(fonte, tipo)


# --- protecao SSRF -------------------------------------------------------------------

@pytest.fixture
def dns(monkeypatch):
    """Dubla o DNS: hostname -> IP. Hostname fora da tabela nao resolve."""
    tabela = {"exemplo.org": "93.184.215.14", "interno.exemplo.org": "192.168.0.10",
              "www.camara.leg.br": "192.168.0.20"}

    def getaddrinfo(host, porta, *a, **k):
        if host not in tabela:
            raise socket.gaierror("nao resolve")
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (tabela[host], 0))]

    monkeypatch.setattr(ex.socket, "getaddrinfo", getaddrinfo)
    monkeypatch.delenv("EXTRACTOR_TRUSTED_DOMAINS", raising=False)
    return tabela


@pytest.mark.parametrize("url", [
    "ftp://exemplo.org/a", "file:///etc/passwd", "http://localhost/admin", "http://metadata.google.internal/",
    "http://169.254.169.254/latest/meta-data/", "http://127.0.0.1:8080/", "http://192.168.0.1/",
    "http://[::1]/", "http://interno.exemplo.org/", "http://nao-resolve.exemplo/", "http:///sem-host",
])
def test_url_perigosa_e_bloqueada(dns, url):
    with pytest.raises(ValueError):
        ex._validate_url(url)


def test_url_publica_passa(dns):
    assert ex._validate_url("  https://exemplo.org/pagina ") == "https://exemplo.org/pagina"


def test_dominio_confiavel_passa_mesmo_com_ip_privado(dns, monkeypatch):
    assert ex._validate_url("https://www.camara.leg.br/x") == "https://www.camara.leg.br/x"   # DNS interno
    monkeypatch.setenv("EXTRACTOR_TRUSTED_DOMAINS", "exemplo.org")   # troca a lista: camara sai
    with pytest.raises(ValueError):
        ex._validate_url("https://www.camara.leg.br/x")
    assert ex._validate_url("https://interno.exemplo.org/") == "https://interno.exemplo.org/"


# --- URL (requests dublado) ------------------------------------------------------------

class _Resp:
    def __init__(self, corpo=b"", status=200, location=None, headers=None, encoding="utf-8"):
        self.corpo, self.status_code, self.encoding = corpo, status, encoding
        self.headers = dict(headers or {})
        if location:
            self.headers["Location"] = location
        self.is_redirect = location is not None

    def close(self):
        pass

    def raise_for_status(self):
        if self.status_code >= 400:
            import requests
            erro = requests.exceptions.HTTPError(f"{self.status_code}")
            erro.response = self
            raise erro

    def iter_content(self, chunk_size):
        for i in range(0, len(self.corpo), chunk_size):
            yield self.corpo[i:i + chunk_size]


@pytest.fixture
def web(monkeypatch, dns):
    """url -> _Resp; guarda as URLs pedidas."""
    paginas, pedidas = {}, []

    def get(url, **k):
        assert k["allow_redirects"] is False and k["stream"] is True
        pedidas.append(url)
        return paginas[url]

    monkeypatch.setattr(ex.requests, "get", get)
    return paginas, pedidas


def test_url_devolve_texto_sem_script_e_estilo(web):
    paginas, _ = web
    paginas["https://exemplo.org/a"] = _Resp(
        "<html><head><style>x{}</style></head><body><script>alert(1)</script><h1>Título</h1><p>Corpo</p>"
        "</body></html>".encode("utf-8"))
    assert extracao_texto.extract_text("https://exemplo.org/a", "url") == "Título\nCorpo"


def test_redirect_absoluto_e_relativo_com_barra_sao_seguidos(web):
    paginas, pedidas = web
    paginas["https://exemplo.org/a"] = _Resp(status=302, location="/b")
    paginas["https://exemplo.org/b"] = _Resp(status=301, location="https://exemplo.org/c")
    paginas["https://exemplo.org/c"] = _Resp(b"<body>fim</body>")
    assert ex.extract_from_url("https://exemplo.org/a") == "fim"
    assert pedidas == ["https://exemplo.org/a", "https://exemplo.org/b", "https://exemplo.org/c"]


def test_redirect_para_endereco_interno_e_bloqueado(web):
    paginas, pedidas = web
    paginas["https://exemplo.org/a"] = _Resp(status=302, location="http://169.254.169.254/latest/meta-data/")
    with pytest.raises(ValueError, match="bloqueado"):
        ex.extract_from_url("https://exemplo.org/a")
    assert pedidas == ["https://exemplo.org/a"]   # o destino interno nunca foi pedido


def test_resposta_grande_demais_e_recusada(web, monkeypatch):
    paginas, _ = web
    monkeypatch.setattr(ex, "_MAX_RESPONSE_SIZE_BYTES", 100)
    paginas["https://exemplo.org/declara"] = _Resp(b"<body>x</body>", headers={"Content-Length": "5000"})
    paginas["https://exemplo.org/mente"] = _Resp(b"<body>" + b"x" * 500 + b"</body>")
    with pytest.raises(RuntimeError, match="limite"):
        ex.extract_from_url("https://exemplo.org/declara")
    with pytest.raises(RuntimeError, match="limite"):
        ex.extract_from_url("https://exemplo.org/mente")


def test_erro_http_vira_runtime_error(web):
    paginas, _ = web
    paginas["https://exemplo.org/404"] = _Resp(status=404)
    with pytest.raises(RuntimeError, match="404"):
        ex.extract_from_url("https://exemplo.org/404")


@pytest.mark.xfail(strict=True, reason="1.0.0: redirect relativo sem barra ('b.html') nao e resolvido contra a URL "
                                       "atual e cai como URL invalida (ValueError). Pendente, ver CHANGELOG.")
def test_redirect_relativo_sem_barra(web):
    paginas, _ = web
    paginas["https://exemplo.org/dir/a.html"] = _Resp(status=302, location="b.html")
    paginas["https://exemplo.org/dir/b.html"] = _Resp(b"<body>ok</body>")
    assert ex.extract_from_url("https://exemplo.org/dir/a.html") == "ok"


def test_versao():
    assert extracao_texto.__version__.count(".") == 2
