# -*- coding: utf-8 -*-
"""
fix_col_widths.py - Adjusts column widths in all .xlsx files inside
docs_operacionais/ based on content length, then recalculates row heights.
"""

import math
import os
import openpyxl
from openpyxl.utils import get_column_letter

CHARS_PER_WIDTH_UNIT = 1.1
LINE_HEIGHT_PT       = 15.0
PADDING_PT           = 5.0
MIN_HEIGHT_PT        = 22.0
TITLE_MIN_HEIGHT_PT  = 40.0
DEFAULT_COL_WIDTH    = 8.43
COL_A_WIDTH          = 5.0
MAX_COL_WIDTH        = 90.0

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(BASE_DIR, "docs_operacionais")
NEWLINE = chr(10)

def get_merged_range_for_cell(ws, row, col):
    """Return the MergedCellRange that contains (row, col), or None."""
    for mr in ws.merged_cells.ranges:
        if (mr.min_row <= row <= mr.max_row and
                mr.min_col <= col <= mr.max_col):
            return mr
    return None


def effective_column_width(ws, col, merged_range=None):
    """Return effective width for text in a cell (summing merged cols)."""
    if merged_range is not None:
        total = 0.0
        for c in range(merged_range.min_col, merged_range.max_col + 1):
            letter = get_column_letter(c)
            dim = ws.column_dimensions.get(letter)
            w = dim.width if (dim and dim.width) else DEFAULT_COL_WIDTH
            total += w
        return total
    else:
        letter = get_column_letter(col)
        dim = ws.column_dimensions.get(letter)
        w = dim.width if (dim and dim.width) else DEFAULT_COL_WIDTH
        return w


def max_single_line_length(text):
    """Return length of longest single line in text."""
    if not text:
        return 0
    text = str(text)
    return max(len(line) for line in text.split(NEWLINE))


def compute_ideal_width(max_len, current_width):
    """Return new ideal width based on max content length."""
    if max_len > 100:
        ideal = max(current_width, 55.0)
        return min(ideal, MAX_COL_WIDTH)
    elif max_len > 50:
        ideal = max(current_width, 40.0)
        return min(ideal, MAX_COL_WIDTH)
    elif max_len > 30:
        ideal = max(current_width, 30.0)
        return min(ideal, MAX_COL_WIDTH)
    elif max_len > 15:
        ideal = max(current_width, 20.0)
        return min(ideal, MAX_COL_WIDTH)
    else:
        return current_width


def estimate_lines(text, col_width):
    """Estimate visual lines for text given column width."""
    if not text:
        return 1
    text = str(text)
    chars_per_line = max(col_width * CHARS_PER_WIDTH_UNIT, 1)
    total_lines = 0
    for segment in text.split(NEWLINE):
        seg_len = len(segment)
        if seg_len == 0:
            total_lines += 1
        else:
            total_lines += math.ceil(seg_len / chars_per_line)
    return total_lines


def required_height(num_lines):
    """Convert a line count into a point height with padding."""
    return num_lines * LINE_HEIGHT_PT + PADDING_PT


def adjust_column_widths(ws):
    """Adjust column widths on a single worksheet."""
    width_changes = []
    max_content_len = {}

    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            value = cell.value
            if value is None:
                continue
            mr = get_merged_range_for_cell(ws, row, col)
            if mr is not None:
                if row != mr.min_row or col != mr.min_col:
                    continue
                if mr.max_col > mr.min_col:
                    continue
            line_len = max_single_line_length(value)
            if col not in max_content_len or line_len > max_content_len[col]:
                max_content_len[col] = line_len

    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        if col == 1:
            dim = ws.column_dimensions.get(letter)
            current_w = dim.width if (dim and dim.width) else DEFAULT_COL_WIDTH
            if abs(current_w - COL_A_WIDTH) >= 0.5:
                width_changes.append((letter, current_w, COL_A_WIDTH))
                ws.column_dimensions[letter].width = COL_A_WIDTH
            continue

        max_len = max_content_len.get(col, 0)
        if max_len == 0:
            continue

        dim = ws.column_dimensions.get(letter)
        current_w = dim.width if (dim and dim.width) else DEFAULT_COL_WIDTH

        new_w = compute_ideal_width(max_len, current_w)
        new_w = round(new_w, 2)

        if abs(new_w - current_w) >= 0.5:
            width_changes.append((letter, round(current_w, 2), new_w))
            ws.column_dimensions[letter].width = new_w

    return width_changes


def recalculate_row_heights(ws):
    """Recalculate row heights based on current column widths."""
    height_changes = []

    for row in range(1, ws.max_row + 1):
        max_needed = 0.0

        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            value = cell.value
            if value is None:
                continue
            mr = get_merged_range_for_cell(ws, row, col)
            if mr is not None:
                if row != mr.min_row or col != mr.min_col:
                    continue
            col_w = effective_column_width(ws, col, mr)
            lines_count = estimate_lines(value, col_w)
            h = required_height(lines_count)
            if h > max_needed:
                max_needed = h

        if max_needed == 0.0:
            continue

        current_height = ws.row_dimensions[row].height

        if row == 1:
            floor = TITLE_MIN_HEIGHT_PT
        else:
            floor = MIN_HEIGHT_PT

        if current_height is not None:
            floor = max(floor, current_height)

        new_height = max(max_needed, floor)
        new_height = round(new_height, 1)

        if current_height is None or abs(new_height - current_height) >= 0.5:
            old_display = current_height if current_height is not None else "auto"
            height_changes.append((row, old_display, new_height))
            ws.row_dimensions[row].height = new_height

    return height_changes


def process_workbook(filepath):
    """Open workbook, adjust widths and heights, save."""
    wb = openpyxl.load_workbook(filepath)
    summary = {}

    for ws in wb.worksheets:
        w_changes = adjust_column_widths(ws)
        h_changes = recalculate_row_heights(ws)

        if w_changes or h_changes:
            summary[ws.title] = {
                "width_changes": w_changes,
                "height_changes": h_changes,
            }

    wb.save(filepath)
    wb.close()
    return summary


def main():
    if not os.path.isdir(DOCS_DIR):
        print(f"ERROR: directory not found: {DOCS_DIR}")
        return

    xlsx_files = sorted(
        f for f in os.listdir(DOCS_DIR) if f.endswith(".xlsx")
    )

    if not xlsx_files:
        print("No .xlsx files found in docs_operacionais/.")
        return

    print(f"Found {len(xlsx_files)} .xlsx file(s) in docs_operacionais/")
    print()
    total_width_changes = 0
    total_height_changes = 0

    for fname in xlsx_files:
        fpath = os.path.join(DOCS_DIR, fname)
        print(f"Processing: {fname}")
        try:
            summary = process_workbook(fpath)
            if summary:
                file_w = 0
                file_h = 0
                for sheet_name, data in summary.items():
                    wc = data["width_changes"]
                    hc = data["height_changes"]
                    if wc:
                        print(f"  Sheet [{sheet_name}] -- {len(wc)} column width(s) adjusted:")
                        for (letter, old_w, new_w) in wc:
                            print(f"    Col {letter}: {old_w} -> {new_w}")
                        file_w += len(wc)
                    if hc:
                        print(f"  Sheet [{sheet_name}] -- {len(hc)} row height(s) recalculated:")
                        for (r, old_h, new_h) in hc:
                            print(f"    Row {r:>3}: {str(old_h):>6} -> {new_h:>6}")
                        file_h += len(hc)
                total_width_changes += file_w
                total_height_changes += file_h
                print(f"  Subtotal: {file_w} width change(s), {file_h} height change(s)")
                print()
            else:
                print("  No changes needed.")
                print()
        except Exception as e:
            import traceback
            print(f"  ERROR: {e}")
            traceback.print_exc()
            print()

    print("=" * 60)
    print(f"Done. Total column width adjustments : {total_width_changes}")
    print(f"      Total row height recalculations: {total_height_changes}")


if __name__ == "__main__":
    main()
