# -*- coding: utf-8 -*-
"""
docs_operacionais_helpers.py
Funções auxiliares e constantes para geração dos documentos operacionais.
"""
import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# ============================================================
# DIRETÓRIO DE SAÍDA
# ============================================================
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs_operacionais")

FOOTER_TEXT = "Portaria 227/2025 — Política de Governança de IA — Câmara dos Deputados"

# ============================================================
# CORES
# ============================================================
DARK_BLUE = "1F4E79"
MEDIUM_BLUE = "2E75B6"
LIGHT_BLUE = "D6E4F0"
WHITE = "FFFFFF"
LIGHT_GRAY = "F2F2F2"
LIGHT_YELLOW = "FFFDE7"
LIGHT_RED = "FDE0DC"
LIGHT_GREEN = "E8F5E9"

# ============================================================
# FONTES
# ============================================================
def font_title():
    return Font(name='Arial', size=14, bold=True, color=WHITE)

def font_subtitle():
    return Font(name='Arial', size=11, bold=True, color=WHITE)

def font_section():
    return Font(name='Arial', size=11, bold=True, color=DARK_BLUE)

def font_label():
    return Font(name='Arial', size=10, bold=True)

def font_normal():
    return Font(name='Arial', size=10)

def font_small():
    return Font(name='Arial', size=9, italic=True, color="666666")

def font_header():
    return Font(name='Arial', size=10, bold=True, color=WHITE)

# ============================================================
# PREENCHIMENTOS
# ============================================================
def fill_dark():
    return PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type='solid')

def fill_medium():
    return PatternFill(start_color=MEDIUM_BLUE, end_color=MEDIUM_BLUE, fill_type='solid')

def fill_light():
    return PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type='solid')

def fill_gray():
    return PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type='solid')

def fill_yellow():
    return PatternFill(start_color=LIGHT_YELLOW, end_color=LIGHT_YELLOW, fill_type='solid')

def fill_red():
    return PatternFill(start_color=LIGHT_RED, end_color=LIGHT_RED, fill_type='solid')

def fill_green():
    return PatternFill(start_color=LIGHT_GREEN, end_color=LIGHT_GREEN, fill_type='solid')

# ============================================================
# ALINHAMENTOS
# ============================================================
ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
ALIGN_TOP = Alignment(horizontal='left', vertical='top', wrap_text=True)

# ============================================================
# BORDAS
# ============================================================
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

# ============================================================
# FUNÇÕES AUXILIARES — XLSX
# ============================================================

def ensure_output_dir():
    os.makedirs(OUTPUT_DIR, exist_ok=True)


def new_workbook():
    """Cria workbook limpo (remove sheet padrão)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


def add_sheet(wb, title):
    """Adiciona sheet com configurações padrão."""
    ws = wb.create_sheet(title=title)
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    return ws


def write_title_row(ws, row, title, num_cols, height=36):
    """Escreve linha de título mesclada com fundo escuro."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_cols)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = font_title()
    cell.fill = fill_dark()
    cell.alignment = ALIGN_CENTER
    cell.border = THIN_BORDER
    ws.row_dimensions[row].height = height
    for c in range(2, num_cols + 1):
        ws.cell(row=row, column=c).fill = fill_dark()
        ws.cell(row=row, column=c).border = THIN_BORDER


def write_subtitle_row(ws, row, text, num_cols, height=24):
    """Escreve linha de subtítulo mesclada com fundo médio."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_cols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = font_subtitle()
    cell.fill = fill_medium()
    cell.alignment = ALIGN_CENTER
    cell.border = THIN_BORDER
    ws.row_dimensions[row].height = height
    for c in range(2, num_cols + 1):
        ws.cell(row=row, column=c).fill = fill_medium()
        ws.cell(row=row, column=c).border = THIN_BORDER


def write_section_header(ws, row, text, num_cols, height=22):
    """Escreve cabeçalho de seção mesclado com fundo azul claro."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_cols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = font_section()
    cell.fill = fill_light()
    cell.alignment = ALIGN_LEFT
    cell.border = THIN_BORDER
    ws.row_dimensions[row].height = height
    for c in range(2, num_cols + 1):
        ws.cell(row=row, column=c).fill = fill_light()
        ws.cell(row=row, column=c).border = THIN_BORDER


def write_col_headers(ws, row, headers, height=30):
    """Escreve cabeçalhos de coluna com fundo médio e texto branco."""
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = font_header()
        cell.fill = fill_medium()
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = height


def write_form_field(ws, row, label, label_col, value_col, value_end_col,
                     ref_artigo=None, is_required=False, height=28):
    """Escreve campo de formulário: label cinza + campo amarelo para preenchimento."""
    label_text = f"{'* ' if is_required else ''}{label}"
    if ref_artigo:
        label_text += f"  [{ref_artigo}]"

    cell_l = ws.cell(row=row, column=label_col, value=label_text)
    cell_l.font = font_label()
    cell_l.fill = fill_gray()
    cell_l.alignment = ALIGN_LEFT
    cell_l.border = THIN_BORDER

    if value_end_col > value_col:
        ws.merge_cells(start_row=row, start_column=value_col,
                       end_row=row, end_column=value_end_col)
    cell_v = ws.cell(row=row, column=value_col)
    cell_v.font = font_normal()
    cell_v.fill = fill_yellow()
    cell_v.alignment = ALIGN_LEFT
    cell_v.border = THIN_BORDER
    for c in range(value_col + 1, value_end_col + 1):
        ws.cell(row=row, column=c).border = THIN_BORDER
        ws.cell(row=row, column=c).fill = fill_yellow()

    ws.row_dimensions[row].height = height


def write_data_row(ws, row, values, height=22, fills=None):
    """Escreve linha de dados com borda fina."""
    for i, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = font_normal()
        cell.alignment = ALIGN_LEFT
        cell.border = THIN_BORDER
        if fills and i <= len(fills) and fills[i - 1]:
            cell.fill = fills[i - 1]
    ws.row_dimensions[row].height = height


def write_checklist_row(ws, row, item_num, item_text, ref_artigo, num_cols, height=28):
    """Escreve linha de checklist: #, Item, Artigo, Status, Evidência, Obs."""
    ws.cell(row=row, column=1, value=item_num).font = font_normal()
    ws.cell(row=row, column=1).alignment = ALIGN_CENTER
    ws.cell(row=row, column=2, value=item_text).font = font_normal()
    ws.cell(row=row, column=2).alignment = ALIGN_LEFT
    ws.cell(row=row, column=3, value=ref_artigo).font = font_small()
    ws.cell(row=row, column=3).alignment = ALIGN_CENTER
    for c in range(4, num_cols + 1):
        ws.cell(row=row, column=c).font = font_normal()
        ws.cell(row=row, column=c).fill = fill_yellow()
        ws.cell(row=row, column=c).alignment = ALIGN_LEFT
    for c in range(1, num_cols + 1):
        ws.cell(row=row, column=c).border = THIN_BORDER
    ws.row_dimensions[row].height = height


def add_dropdown(ws, cell_range, options):
    """Adiciona validação de dados (drop-down) a um intervalo."""
    formula = '"' + ','.join(options) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Selecione uma opção da lista"
    dv.errorTitle = "Valor inválido"
    dv.prompt = "Selecione uma opção"
    dv.promptTitle = "Opções"
    ws.add_data_validation(dv)
    dv.add(cell_range)


def write_footer(ws, row, num_cols, height=20):
    """Escreve rodapé com referência normativa."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_cols)
    cell = ws.cell(row=row, column=1, value=FOOTER_TEXT)
    cell.font = font_small()
    cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = height


def set_col_widths(ws, widths):
    """Define larguras de colunas. widths = lista de floats."""
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def save_xlsx(wb, filename):
    """Salva workbook no diretório de saída."""
    ensure_output_dir()
    path = os.path.join(OUTPUT_DIR, filename)
    wb.save(path)
    print(f"  Salvo: {path}")
    return path


# ============================================================
# FUNÇÕES AUXILIARES — DOCX
# ============================================================

def new_document(title, subtitle=None):
    """Cria documento Word com cabeçalho institucional."""
    doc = Document()

    # Configurar estilos padrão
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    # Cabeçalho institucional
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CÂMARA DOS DEPUTADOS")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run(subtitle)
        run2.font.size = Pt(11)
        run2.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    # Título do documento
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.add_run(title)
    run3.bold = True
    run3.font.size = Pt(13)
    run3.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    # Linha separadora
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run4 = p4.add_run("—" * 60)
    run4.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    run4.font.size = Pt(10)

    return doc


def add_docx_section(doc, title):
    """Adiciona seção com título formatado."""
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(12)


def add_docx_field(doc, label, ref_artigo=None, lines=2):
    """Adiciona campo preenchível no documento Word."""
    text = f"{label}"
    if ref_artigo:
        text += f"  [{ref_artigo}]"
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10)

    for _ in range(lines):
        p2 = doc.add_paragraph()
        run2 = p2.add_run("_" * 80)
        run2.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        run2.font.size = Pt(10)


def add_docx_paragraph(doc, text, bold=False, italic=False, size=11):
    """Adiciona parágrafo formatado."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return p


def add_docx_table(doc, headers, rows, col_widths=None):
    """Adiciona tabela formatada."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Cabeçalhos
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.name = 'Arial'
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        from docx.oxml.ns import qn
        shading = cell._element.get_or_add_tcPr()
        shading_elm = openpyxl.xml.functions = None  # placeholder
        # Usar XML direto para shading
        from lxml import etree
        ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        tc_pr = cell._element.get_or_add_tcPr()
        shd = etree.SubElement(tc_pr, qn('w:shd'))
        shd.set(qn('w:fill'), '2E75B6')
        shd.set(qn('w:val'), 'clear')

    # Dados
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val) if val else ""
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.name = 'Arial'

    return table


def add_signature_block(doc, roles):
    """Adiciona bloco de assinaturas."""
    doc.add_paragraph()  # espaço
    for role in roles:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(24)
        run = p.add_run("_" * 50)
        run.font.size = Pt(10)

        p2 = doc.add_paragraph()
        run2 = p2.add_run(role)
        run2.bold = True
        run2.font.size = Pt(10)

        p3 = doc.add_paragraph()
        run3 = p3.add_run("Data: ____/____/________")
        run3.font.size = Pt(10)


def add_docx_footer(doc):
    """Adiciona rodapé com referência normativa."""
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("—" * 60)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    run.font.size = Pt(9)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(FOOTER_TEXT)
    run2.font.size = Pt(9)
    run2.italic = True
    run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def save_docx(doc, filename):
    """Salva documento Word no diretório de saída."""
    ensure_output_dir()
    path = os.path.join(OUTPUT_DIR, filename)
    doc.save(path)
    print(f"  Salvo: {path}")
    return path
