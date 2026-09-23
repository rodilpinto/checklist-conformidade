# -*- coding: utf-8 -*-
"""
fix_row_heights.py
==================
Adjusts row heights in all .xlsx files inside docs_operacionais/ so that
wrapped text is fully visible.

Algorithm per cell:
  1. Get the effective column width (sum of merged columns if the cell is
     the top-left of a merge range).
  2. Estimate characters per line  = column_width * CHARS_PER_WIDTH_UNIT
  3. Estimate wrapped lines        = ceil(len(text) / chars_per_line)
  4. Account for explicit newlines  (each '\n' adds a line)
  5. Required height                = total_lines * LINE_HEIGHT_PT + PADDING_PT
  6. Row height                     = max(required across all cells in row,
                                          existing explicit height,
                                          MIN_HEIGHT_PT)
  7. Row 1 always gets at least TITLE_MIN_HEIGHT_PT.

Constants are tunable at the top of the script.
"""

import math
import os
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

# -- Tunable constants -------------------------------------------------------
CHARS_PER_WIDTH_UNIT = 1.1   # approx characters fitting in 1 unit of col width
LINE_HEIGHT_PT       = 15.0  # points per line of text
PADDING_PT           = 5.0   # extra padding added to every calculated height
MIN_HEIGHT_PT        = 22.0  # absolute minimum row height
TITLE_MIN_HEIGHT_PT  = 40.0  # minimum height for row 1 (title)
DEFAULT_COL_WIDTH    = 8.43  # openpyxl default when no explicit width is set

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(BASE_DIR, "docs_operacionais")


def get_merged_range_for_cell(ws, row, col):
    """Return the MergedCellRange that contains (row, col), or None."""
    for mr in ws.merged_cells.ranges:
        if (mr.min_row <= row <= mr.max_row and
                mr.min_col <= col <= mr.max_col):
            return mr
    return None


def effective_column_width(ws, col, merged_range=None):
    """
    Return the effective width available for text in a cell.
    For merged cells, sum the widths of all columns in the merge.
    """
    if merged_range is not None:
        total = 0.0
        for c in range(merged_range.min_col, merged_range.max_col + 1):
            letter = get_column_letter(c)
            dim = ws.column_dimensions.get(letter)
            w = dim.width if (dim and dim.width) else DEFAULT_COL_WIDTH
            total += w
        return total
    else:
        letter = get_column_letter(col)
        dim = ws.column_dimensions.get(letter)
        w = dim.width if (dim and dim.width) else DEFAULT_COL_WIDTH
        return w


def estimate_lines(text, col_width):
    """
    Estimate how many visual lines `text` will occupy given column width.
    Accounts for explicit newlines inside the text.
    """
    if not text:
        return 1
    text = str(text)
    chars_per_line = max(col_width * CHARS_PER_WIDTH_UNIT, 1)
    total_lines = 0
    for segment in text.split("\n"):
        seg_len = len(segment)
        if seg_len == 0:
            total_lines += 1
        else:
            total_lines += math.ceil(seg_len / chars_per_line)
    return total_lines


def required_height(num_lines):
    """Convert a line count into a point height with padding."""
    return num_lines * LINE_HEIGHT_PT + PADDING_PT


def process_workbook(filepath):
    """
    Open a workbook, fix row heights on every sheet, save, and return
    a dict of changes:  {sheet_name: [(row, old_height, new_height), ...]}
    """
    wb = openpyxl.load_workbook(filepath)
    changes = {}

    for ws in wb.worksheets:
        sheet_changes = []

        for row in range(1, ws.max_row + 1):
            # Determine the tallest requirement across all cells in this row
            max_needed = 0.0

            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row, col)
                value = cell.value
                if value is None:
                    continue

                # Check if this cell is part of a merged range
                mr = get_merged_range_for_cell(ws, row, col)

                # For merged cells, only process the top-left cell
                if mr is not None:
                    if row != mr.min_row or col != mr.min_col:
                        continue  # skip non-origin merged cells

                col_w = effective_column_width(ws, col, mr)
                lines = estimate_lines(value, col_w)
                h = required_height(lines)
                if h > max_needed:
                    max_needed = h

            if max_needed == 0.0:
                # Row has no content -- skip
                continue

            # Current (explicit) height; None means "auto" which we treat as default
            current_height = ws.row_dimensions[row].height

            # Determine minimum for this row
            if row == 1:
                floor = TITLE_MIN_HEIGHT_PT
            else:
                floor = MIN_HEIGHT_PT

            # If there is already an explicit height, never shrink below it
            if current_height is not None:
                floor = max(floor, current_height)

            new_height = max(max_needed, floor)
            # Round to one decimal
            new_height = round(new_height, 1)

            # Only record a change if the height actually differs
            if current_height is None or abs(new_height - current_height) >= 0.5:
                old_display = current_height if current_height is not None else "auto"
                sheet_changes.append((row, old_display, new_height))
                ws.row_dimensions[row].height = new_height

        if sheet_changes:
            changes[ws.title] = sheet_changes

    wb.save(filepath)
    wb.close()
    return changes


def main():
    if not os.path.isdir(DOCS_DIR):
        print(f"ERROR: directory not found: {DOCS_DIR}")
        return

    xlsx_files = sorted(
        f for f in os.listdir(DOCS_DIR) if f.endswith(".xlsx")
    )

    if not xlsx_files:
        print("No .xlsx files found in docs_operacionais/.")
        return

    print(f"Found {len(xlsx_files)} .xlsx files in docs_operacionais/\n")
    total_changes = 0

    for fname in xlsx_files:
        fpath = os.path.join(DOCS_DIR, fname)
        print(f"Processing: {fname}")
        try:
            changes = process_workbook(fpath)
            if changes:
                file_count = 0
                for sheet_name, rows in changes.items():
                    print(f"  Sheet '{sheet_name}': {len(rows)} row(s) adjusted")
                    for (r, old_h, new_h) in rows:
                        print(f"    Row {r:>3}: {str(old_h):>6} -> {new_h:>6}")
                    file_count += len(rows)
                total_changes += file_count
                print(f"  Subtotal: {file_count} change(s)\n")
            else:
                print("  No changes needed.\n")
        except Exception as e:
            print(f"  ERROR: {e}\n")

    print("=" * 60)
    print(f"Done. Total row height adjustments: {total_changes}")


if __name__ == "__main__":
    main()
