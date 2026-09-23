import openpyxl
from copy import copy
import shutil

import os
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

shutil.copy2(
    os.path.join(_ROOT, 'checklists', 'historico', 'Checklist_Portaria_227_2025_IA_v1.05.xlsx'),
    os.path.join(_ROOT, 'checklists', 'Checklist_Portaria_227_2025_IA_v1.06.xlsx')
)

wb = openpyxl.load_workbook(os.path.join(_ROOT, 'checklists', 'Checklist_Portaria_227_2025_IA_v1.06.xlsx'))

# Find sheets
acoes_sheet = None
curso_sheet = None
for s in wb.sheetnames:
    if 'Atores' in s:
        acoes_sheet = wb[s]
    if 'Curso' in s or 'SECIN' in s:
        curso_sheet = wb[s]

print(f"Ações sheet: {acoes_sheet.title}")
print(f"Curso sheet: {curso_sheet.title}")

# ============================================================
# 1. Ajustes na aba Ações (J10, J127, H padronização)
# ============================================================
acoes_sheet['J10'] = (
    "Nota: o texto publicado no Boletim Administrativo traz "
    "'o seu o desenvolvimento' (provável erro de digitação no original); "
    "reproduzido fielmente nesta checklist."
)
acoes_sheet['J127'] = (
    "Nota: o texto publicado no Boletim Administrativo omite os "
    "dois-pontos após 'CETIA' (presente nos Arts. 29, 30, 31 e 32 "
    "análogos); adicionados nesta checklist por uniformidade."
)

h_fixes = {97, 98, 104, 105, 106}
for row in h_fixes:
    cell = acoes_sheet[f'H{row}']
    if cell.value and 'Encarregado de Dados' in str(cell.value) and 'Pessoais' not in str(cell.value):
        cell.value = str(cell.value).replace('Encarregado de Dados', 'Encarregado de Dados Pessoais')

print("Ações sheet adjustments done")

# ============================================================
# 2. Fix F3 no Curso sheet
# ============================================================
curso_sheet['F3'] = (
    "Módulos 1-6 (Fundamentos, Prompt, Uso Responsável, "
    "Assistentes, Contexto, Agentes) + todos os práticos"
)
print("F3 fixed")

# ============================================================
# 3. Style copy helper
# ============================================================
def copy_style(src, tgt):
    if src.font:
        tgt.font = copy(src.font)
    if src.alignment:
        tgt.alignment = copy(src.alignment)
    if src.fill:
        try:
            tgt.fill = copy(src.fill)
        except Exception:
            pass
    if src.border:
        tgt.border = copy(src.border)
    if src.number_format:
        tgt.number_format = src.number_format

# ============================================================
# 4. New actions
# ============================================================
GN = "Gestor de Negócio\n(SECIN)"
GP = "Gerente de Projeto\n(Capacitação)"

new_actions = [
    # === PRÉ-CONTRATAÇÃO: Lacunas ===
    (28, GN, "Pré-Contratação",
     "Verificar conformidade do projeto de capacitação com a Política de Governança de Dados, a Política de Segurança da Informação e Cibernética e a Política de Proteção de Dados Pessoais",
     "Art. 13, Pú",
     "Todos",
     "Curso inicia sem verificação de conformidade com as três políticas complementares à Portaria 227, gerando risco de violação cruzada em matéria de dados, segurança e privacidade",
     "Alto",
     "Aplicar checklist de conformidade das 3 políticas antes do início do curso; registrar resultado e obter ciência das áreas responsáveis",
     "Checklist de conformidade com as 3 políticas preenchido e aprovado",
     "Gestor de Negócio / CGSIC / Encarregado"),

    (29, GN, "Pré-Contratação",
     "Submeter a contratação do curso à deliberação da Ditec, por se tratar de solução que envolve uso de sistemas de IA",
     "Art. 29, III",
     "Todos",
     "Contratação de capacitação com IA formalizada sem que a Ditec tenha deliberado, violando o Art. 29, III",
     "Alto",
     "Incluir a contratação na pauta de deliberação da Ditec antes da assinatura do contrato; obter registro formal de aprovação",
     "Deliberação da Ditec registrada; parecer técnico favorável",
     "Gestor de Negócio / Ditec"),

    (30, GN, "Pré-Contratação",
     "Garantir que a avaliação de riscos submetida ao CETIA inclua triagem de segurança da informação para consulta ao CGSIC, dado o envolvimento de dados sigilosos de auditoria",
     "Art. 24, § 3º",
     "Todos",
     "CETIA emite parecer sobre o curso/projeto sem consultar o CGSIC, ignorando os riscos de segurança inerentes a dados de auditoria interna",
     "Alto",
     "Incluir questionário de triagem de segurança no formulário de avaliação de riscos; sinalizar ao CETIA a necessidade de consulta ao CGSIC",
     "Registro de consulta ao CGSIC; parecer do CGSIC anexado ao parecer CETIA",
     "Gestor de Negócio / CETIA / CGSIC"),

    (31, GP, "Pré-Contratação",
     "Solicitar ao Gestor de Dados autorização formal para acesso e tratamento de dados nos exercícios do curso; para dados pessoais, solicitar também ao Encarregado",
     "Art. 32, I",
     "Módulos 11-13 (Documentos, Tabulação, Dados)",
     "Exercícios do curso utilizam dados da Câmara sem autorização formal do Gestor de Dados, configurando acesso não autorizado",
     "Alto",
     "Preencher formulário de autorização antes da preparação do material; obter aprovação do Gestor de Dados (e Encarregado, se dados pessoais)",
     "Formulário de autorização aprovado pelo Gestor de Dados (e Encarregado, se aplicável)",
     "Gerente de Projeto / Gestor de Dados"),

    (32, GP, "Pré-Contratação",
     "Verificar onde está hospedado o Ambiente Virtual Exclusivo do curso e garantir que nenhum dado restrito ou pessoal da Câmara será carregado na plataforma; incluir cláusula contratual sobre tratamento e exclusão dos dados",
     "Arts. 8º, 19",
     "Todos (plataforma transversal)",
     "Documentos e dados da Câmara armazenados na plataforma virtual da contratada sem controle de segurança, jurisdição ou exclusão pós-curso",
     "Alto",
     "Exigir informações sobre hospedagem; incluir cláusula contratual de vedação de dados restritos, tratamento e exclusão pós-curso; vedar upload de documentos reais",
     "Cláusula contratual sobre dados; declaração de hospedagem; procedimento de exclusão pós-curso",
     "Gerente de Projeto / Ditec / CGSIC"),

    # === DURANTE O CURSO: Lacuna ===
    (33, GP, "Durante o Curso",
     "Garantir que o Módulo 16 (Comunicação Eficiente) inclua orientação sobre revisão obrigatória antes de publicar resultados de auditoria gerados com auxílio de IA",
     "Art. 14, caput",
     "Módulo 16 (Comunicação Eficiente)",
     "Participantes criam apresentações ou sites com resultados de auditoria gerados por IA e publicam sem revisão adequada",
     "Alto",
     "Incluir orientação no módulo sobre revisão obrigatória; reforçar que qualquer publicação de resultados deve ser validada pelo auditor responsável",
     "Material do Módulo 16 com orientação sobre revisão; exercício prático de validação",
     "Gerente de Projeto / Instrutor"),

    # === PÓS-CURSO: Lacunas ===
    (34, GN, "Pós-Curso",
     "Orientar auditores que conclusões de auditoria apoiadas por IA devem ser verificadas quanto a vieses e não podem gerar resultados discriminatórios contra entidades auditadas",
     "Art. 15, I-III",
     "Módulos 7-8 (Riscos, Planejamento), 11-13 (Documentos, Dados)",
     "IA direciona achados de auditoria com viés contra determinados fornecedores, unidades ou servidores, sem que o auditor perceba a distorção",
     "Alto",
     "Incluir verificação de vieses no procedimento de revisão de relatórios; orientar auditores sobre risco de viés algorítmico em análises de dados",
     "Procedimento de verificação de vieses documentado; orientação publicada aos auditores",
     "Gestor de Negócio"),

    (35, GN, "Pós-Curso",
     "Definir política de transparência: relatórios e produtos de auditoria elaborados com auxílio de IA devem indicar essa circunstância aos destinatários",
     "Art. 12",
     "Módulos 4, 6, 14 (Assistentes, Agentes, Relatoria)",
     "Relatórios de auditoria elaborados com auxílio de IA publicados sem informar aos destinatários, comprometendo transparência e confiança no controle interno",
     "Médio",
     "Incluir campo obrigatório nos relatórios de auditoria indicando se IA foi utilizada; definir política de disclosure da SECIN",
     "Campo de disclosure em template de relatório; política de transparência publicada",
     "Gestor de Negócio"),

    (36, GN, "Pós-Curso",
     "Incluir os fluxos de dados da SECIN que passem por ferramentas de IA no monitoramento contínuo de segurança da informação e cibernética",
     "Art. 27, § 3º",
     "Todos (uso contínuo)",
     "Dados de auditoria submetidos a ferramentas de IA sem monitoramento de segurança, vulneráveis a ataques que comprometam integridade e privacidade",
     "Alto",
     "Mapear fluxos de dados SECIN-IA; incluir no escopo do SIEM/SOC; definir alertas de integridade; testar periodicamente",
     "Mapa de fluxos de dados; inclusão no monitoramento CGSIC; relatórios de testes",
     "Gestor de Negócio / CGSIC / Ditec"),

    (37, GP, "Pós-Curso",
     "Submeter à Ditec o conjunto de prompts fornecidos no curso para validação antes da adoção nos processos de trabalho de auditoria",
     "Art. 29, II",
     "Módulo 2 (Engenharia de Prompt) + biblioteca de prompts do curso",
     "Prompts fornecidos pelo instrutor adotados nos processos de auditoria sem validação pela Ditec, podendo conter instruções inadequadas ou inseguras",
     "Médio",
     "Catalogar todos os prompts recebidos; submeter à Ditec para análise de conformidade com diretrizes técnicas; adotar apenas os aprovados",
     "Catálogo de prompts submetido; aprovação da Ditec registrada",
     "Gerente de Projeto / Ditec"),

    # === DESENVOLVIMENTO DE MÓDULOS DE AUDITORIA COM IA ===
    (38, GN, "Desenv. Módulos Auditoria",
     "Submeter a demanda de desenvolvimento de módulos de auditoria com IA à avaliação de riscos da Ditec e ao parecer do CETIA, como novo projeto de sistema de IA",
     "Art. 24, caput e §§ 1º-3º",
     "Originado dos Módulos 7-8 e 14-15 (Riscos, Planejamento, Relatoria, Monitoramento)",
     "Desenvolvimento de módulos de auditoria com IA iniciado como extensão informal do curso, sem avaliação formal de riscos e sem parecer do CETIA",
     "Crítico",
     "Tratar o desenvolvimento como projeto formal de IA; submeter avaliação de riscos à Ditec; obter parecer do CETIA antes de iniciar",
     "Avaliação de riscos submetida; parecer CETIA obtido; aprovação formal do projeto",
     "Gestor de Negócio / Ditec / CETIA"),

    (39, GN, "Desenv. Módulos Auditoria",
     "Formalizar o projeto junto à Ditec seguindo o processo de desenvolvimento e sustentação de sistemas de IA, com gates éticos em cada fase do ciclo de vida",
     "Art. 29, I",
     "Todos (evolução pós-curso)",
     "Módulos de auditoria com IA desenvolvidos fora do processo formal da Ditec, sem práticas de conformidade com princípios éticos da Portaria 227",
     "Crítico",
     "Registrar projeto no portfólio da Ditec; seguir processo formal com gates éticos; designar equipe de desenvolvimento com participação de negócio e tecnologia",
     "Projeto registrado no portfólio Ditec; equipe designada; processo formal adotado com gates",
     "Gestor de Negócio / Ditec"),

    (40, GN, "Desenv. Módulos Auditoria",
     "Fornecer tempestivamente requisitos de negócio de auditoria que possam interferir no comportamento dos módulos de IA; participar de testes, simulações e validações",
     "Art. 31, VII",
     "Todos (evolução pós-curso)",
     "Módulos de IA desenvolvidos sem requisitos claros de auditoria, gerando resultados desalinhados com necessidades reais e padrões do TCU/CGU",
     "Alto",
     "Designar auditores-chave para levantamento de requisitos; participar de sprints de validação; documentar requisitos com priorização",
     "Requisitos de negócio documentados e entregues; registros de participação em testes e validações",
     "Gestor de Negócio / Ditec"),

    (41, GN, "Desenv. Módulos Auditoria",
     "Deliberar, em parceria com Ditec e Gestores de Dados, sobre a qualidade e adequação dos dados de auditoria que alimentarão os módulos de IA",
     "Arts. 31, V; 29, VII",
     "Módulos 11-13 (Documentos, Tabulação, Dados)",
     "Módulos de IA alimentados com dados de auditoria incompletos, desatualizados ou de baixa qualidade, gerando resultados não confiáveis",
     "Alto",
     "Realizar avaliação de qualidade de dados antes do uso; definir critérios mínimos; estabelecer pipeline de atualização e limpeza",
     "Parecer de qualidade de dados; critérios definidos; pipeline documentado",
     "Gestor de Negócio / Ditec / Gestores de Dados"),

    (42, GN, "Desenv. Módulos Auditoria",
     "Garantir que todos os módulos de auditoria com IA mantenham a primazia da decisão humana, com revisão obrigatória do auditor em cada etapa decisória",
     "Art. 14, Pú; Art. 4º, VII",
     "Todos os módulos desenvolvidos",
     "Módulos de IA produzem achados, recomendações ou pareceres adotados automaticamente sem revisão do auditor, comprometendo autonomia humana e qualidade do controle",
     "Crítico",
     "Implementar workflow com etapa obrigatória de revisão humana em cada output; vedar decisões automatizadas sem aprovação do auditor",
     "Workflow com human-in-the-loop documentado e testado; vedação de automação total verificada",
     "Gestor de Negócio / Ditec"),

    (43, GN, "Desenv. Módulos Auditoria",
     "Planejar a supervisão contínua compartilhada (SECIN, Gestores de Dados, Ditec) desde a concepção dos módulos, incluindo testes éticos periódicos",
     "Art. 27, caput e §§ 1º-2º",
     "Todos (operação futura)",
     "Módulos implantados sem plano de supervisão, gerando degradação não detectada de desempenho, vieses emergentes ou problemas éticos",
     "Alto",
     "Definir plano de supervisão desde a fase de concepção; incluir cronograma de testes éticos periódicos; designar responsáveis por aspecto",
     "Plano de supervisão contínua; cronograma de testes; responsáveis designados",
     "Gestor de Negócio / Ditec / Gestores de Dados"),

    (44, GP, "Desenv. Módulos Auditoria",
     "Garantir atuação conjunta da SECIN e da Ditec com atribuições definidas em cada fase do ciclo de vida dos módulos de IA",
     "Art. 26, § 1º",
     "Todos (evolução pós-curso)",
     "Desenvolvimento concentrado apenas na Ditec ou na SECIN, sem visão integrada de negócio e tecnologia, gerando módulos desalinhados",
     "Alto",
     "Publicar matriz de responsabilidades por fase; realizar reuniões conjuntas periódicas; documentar decisões e atribuições",
     "Matriz de responsabilidades publicada; atas de reuniões conjuntas; decisões documentadas",
     "Gerente de Projeto / Gestor de Negócio / Ditec"),

    (45, GP, "Desenv. Módulos Auditoria",
     "Aplicar gestão de riscos de IA em cada fase do ciclo de vida dos módulos, desde a concepção até eventual desativação ou remoção do portfólio",
     "Art. 22, caput",
     "Todos (evolução pós-curso)",
     "Riscos de IA não gerenciados ao longo do desenvolvimento, acumulando vulnerabilidades, problemas éticos e dívida técnica",
     "Alto",
     "Aplicar framework de gestão de riscos da Ditec; revisão de riscos em cada gate de fase; manter e atualizar matriz de riscos do projeto",
     "Matriz de riscos por fase; registros de revisão em cada gate; atualizações documentadas",
     "Gerente de Projeto / Ditec"),

    (46, GP, "Desenv. Módulos Auditoria",
     "Elaborar plano de testes de robustez e segurança; incluir identificação de riscos de incidentes e testes para detectar comportamentos inesperados",
     "Arts. 16; 26, § 3º",
     "Todos os módulos desenvolvidos",
     "Módulos de auditoria com IA implantados sem testes de resiliência, vulneráveis a falhas, erros e ataques de segurança",
     "Crítico",
     "Definir plano de testes (unitários, integração, stress, adversariais); gate de aprovação antes de implantação; testar com dados representativos anonimizados",
     "Plano de testes; relatório de execução; aprovação formal para implantação",
     "Gerente de Projeto / Ditec / CGSIC"),

    (47, GP, "Desenv. Módulos Auditoria",
     "Implementar medidas de identificação e controle de vieses e alucinações que possam distorcer achados de auditoria",
     "Art. 26, § 4º",
     "Todos os módulos desenvolvidos",
     "Módulos de IA geram achados distorcidos por vieses nos dados ou alucinações, comprometendo a credibilidade do controle interno da Câmara",
     "Crítico",
     "Incluir testes de vieses e alucinações no plano de testes; monitorar em produção; definir métricas de qualidade; implementar RAG com bases confiáveis",
     "Relatório de testes de vieses/alucinações; métricas de qualidade; fontes RAG documentadas",
     "Gerente de Projeto / Ditec"),

    (48, GP, "Desenv. Módulos Auditoria",
     "Garantir transparência, explicabilidade e rastreabilidade; implementar mecanismos de verificação resultados-dados, recuperação de dados de treinamento e armazenamento de fontes RAG",
     "Arts. 10, 11, 25 I-III",
     "Todos os módulos desenvolvidos",
     "Módulos de auditoria operam como caixa-preta, impossibilitando explicar ou auditar as decisões e recomendações geradas pela IA",
     "Crítico",
     "Implementar logging de decisões; documentar modelo, features e dados; manter repositório de dados de treinamento versionado; catalogar fontes RAG com snapshots",
     "Logs de decisão acessíveis; documentação do modelo; repositório versionado; catálogo de fontes RAG",
     "Gerente de Projeto / Ditec"),

    (49, GP, "Desenv. Módulos Auditoria",
     "Garantir confidencialidade dos dados de auditoria conforme grau de sigilo; obter autorizações para nuvem privada; vedar processamento em nuvem pública",
     "Arts. 17, 18, 19",
     "Todos os módulos desenvolvidos",
     "Dados sigilosos de auditoria processados por módulos de IA em nuvem pública ou sem controles de acesso adequados ao grau de sigilo",
     "Crítico",
     "Classificar dados por grau de sigilo; implementar controles de acesso; processar em nuvem privada ou infraestrutura própria; obter autorização do Gestor de Dados",
     "Classificação de dados; controles de acesso; autorizações formais; declaração de nuvem privada/infra própria",
     "Gerente de Projeto / CGSIC / Gestor de Dados"),

    (50, GP, "Desenv. Módulos Auditoria",
     "Testar módulos de auditoria para resultados discriminatórios, vieses e perfilamento antes da implantação",
     "Art. 15, I-III",
     "Módulos 7-8 (Riscos, Planejamento), 11-13 (Documentos, Dados)",
     "Módulos de IA geram achados enviesados contra determinadas unidades, fornecedores ou servidores, configurando discriminação nos trabalhos de auditoria",
     "Alto",
     "Realizar testes de fairness com dados representativos; verificar ausência de perfilamento; obter certificação antes da implantação",
     "Relatório de fairness; certificação de ausência de perfilamento; aprovação para implantação",
     "Gerente de Projeto / Ditec / CETIA"),

    (51, GP, "Desenv. Módulos Auditoria",
     "Respeitar propriedade intelectual dos conteúdos nos dados utilizados para treinamento e contexto dos módulos de IA",
     "Art. 20",
     "Todos os módulos desenvolvidos",
     "Dados protegidos por propriedade intelectual ou direitos autorais utilizados em treinamento ou contexto sem autorização",
     "Alto",
     "Verificar licenças de todos os datasets; registrar origem dos dados; obter autorização para material protegido",
     "Registro de licenças; rastreabilidade de origem; autorizações obtidas",
     "Gerente de Projeto / Ditec / Gestor de Negócio"),

    (52, GP, "Desenv. Módulos Auditoria",
     "Se os módulos utilizarem componentes de IA de terceiros (APIs, modelos, agentes), definir requisitos, cláusulas contratuais de governança, monitoramento e procedimento de remoção",
     "Art. 26, § 2º",
     "Módulo 6 (Agentes de IA) + componentes externos",
     "Componentes de IA de terceiros integrados sem cláusulas de governança, dificultando monitoramento, substituição e remoção do portfólio",
     "Alto",
     "Documentar todos os componentes de terceiros; incluir cláusulas de governança de IA em contratos; definir SLAs e procedimento de remoção",
     "Inventário de componentes; cláusulas contratuais; SLAs definidos; procedimento de remoção",
     "Gerente de Projeto / Ditec / Gestor de Negócio"),
]

# ============================================================
# 5. Write new actions
# ============================================================
col_indices = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']
start_row = curso_sheet.max_row + 1

for i, action in enumerate(new_actions):
    row = start_row + i
    for j, col in enumerate(col_indices):
        cell = curso_sheet[f'{col}{row}']
        cell.value = action[j]
        src = curso_sheet[f'{col}2']
        copy_style(src, cell)

print(f"Added {len(new_actions)} new actions (rows {start_row}-{start_row + len(new_actions) - 1})")
print(f"Total actions in Curso sheet: {curso_sheet.max_row - 1}")

_out = os.path.join(_ROOT, 'checklists', 'Checklist_Portaria_227_2025_IA_v1.06.xlsx')
wb.save(_out)
print(f"Saved {_out}")
